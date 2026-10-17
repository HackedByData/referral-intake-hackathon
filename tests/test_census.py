import io
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
from sqlmodel import Session, SQLModel, select

import app.beds.models  # noqa: F401
from app.beds import census, service
from app.beds.models import Assignment, Bed, Site
from app.beds.seed import seed_poc_beds
from app.db import make_engine

REAL_WORKBOOK = Path("C:/Users/devin/UCI/Claude Hackathon/Test Census.xlsx")


def build_workbook() -> bytes:
    """INVENTED census workbook exercising the parser's quirks."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RC Test"
    ws.append(["Bed", "IF#", "FIRST", "LAST", "GENDER", "Contract", "INTAKE DATE", "Exit Date", "PETS", "Total: 3"])
    ws.append(["101-A", "IF2099-1", "Alpha", "T.", "Female", "Cal Optima", datetime(2026, 9, 1), "TBD", "No"])
    ws.append(["101-B", None, "Open"])
    ws.append(["102 (ADA)", None, "Needs Cleaning"])
    ws.append(["103", None, "Closed"])
    ws.append(["Staff", None, "Staff"])
    ws.append(["M1", None, "Bravo", None, "Male", None, "09/05/2026", "10/01/2026"])
    ws.append(["W1", None, None])
    ws.append(["Total Client", 4])
    tw = wb.create_sheet("Twin Test")
    tw.append(["Twin Test census"])
    tw.append(["Room", "IF #", "First Name", "Last Name", "Intake Date", "Notes",
               "Room", "IF #", "First Name", "Last Name", "Intake Date", "Notes"])
    tw.append(["1", "IF2099-2", "Charlie", "X.", datetime(2026, 8, 1), "left note",
               "11", "IF2099-3", "Delta", "Y.", datetime(2026, 8, 2), "right note"])
    tw.append(["2", None, "Open", None, None, None, "12", None, "Open"])
    jk = wb.create_sheet("Junk Test")
    jk.append(["///", "IF#", "First Name", "Last", "Intake Date", "Notes", "Room #", "IF#", "First Name", "Last"])
    jk.append(["201", None, "Open", None, None, None, "301", "IE2099-4", "Echo", "Z."])
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def session():
    engine = make_engine("sqlite://")
    SQLModel.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def parsed():
    return census.parse_workbook(build_workbook())


def by_sheet(sheets):
    return {s.sheet: s for s in sheets}


def test_parse_counts_statuses_and_skips():
    sheets, skipped = parsed()
    sh = by_sheet(sheets)
    assert any(s.startswith("Empty") for s in skipped)
    rc = sh["RC Test"]
    rows = {r.label: r for r in rc.rows}
    assert set(rows) == {"101-A", "101-B", "102 (ADA)", "103", "M1", "W1"}
    assert rows["101-A"].status == "occupied" and rows["101-B"].status == "available"
    assert rows["102 (ADA)"].status == "needs_cleaning"
    assert rows["103"].status == "out_of_service" and "closed" in rows["103"].bed_notes.lower()
    assert any("Staff" in s for s in rc.skipped) and any("Total Client" in s for s in rc.skipped)
    assert rows["102 (ADA)"].wheelchair is True and rows["101-A"].wheelchair is False
    assert rows["M1"].sex_designation == "male" and rows["W1"].sex_designation == "female"
    assert rows["101-A"].sex_designation == "any"


def test_parse_stay_fields_dates_and_extra():
    rc = by_sheet(parsed()[0])["RC Test"]
    rows = {r.label: r for r in rc.rows}
    st = rows["101-A"].stay
    assert st["client_id"] == "IF2099-1" and st["first_name"] == "Alpha" and st["last_name"] == "T."
    assert st["gender"] == "Female" and st["payer"] == "Cal Optima"
    assert st["intake_date"] == date(2026, 9, 1) and st["exit_date"] is None  # TBD
    assert st["extra"]["PETS"] == "No"
    m1 = rows["M1"].stay
    assert m1["client_id"] in (None, "") and m1["first_name"] == "Bravo"
    assert m1["intake_date"] == date(2026, 9, 5) and m1["exit_date"] == date(2026, 10, 1)
    assert rows["101-B"].stay is None


def test_parse_side_by_side_and_junk_headers():
    sh = by_sheet(parsed()[0])
    tw = {r.label: r for r in sh["Twin Test"].rows}
    assert set(tw) == {"1", "2", "11", "12"}
    assert tw["1"].stay["notes"] == "left note" and tw["11"].stay["notes"] == "right note"
    assert tw["11"].stay["first_name"] == "Delta"
    jk = {r.label: r for r in sh["Junk Test"].rows}
    assert set(jk) == {"201", "301"} and jk["301"].stay["client_id"] == "IE2099-4"


def test_import_is_idempotent(session):
    sheets, _ = parsed()
    r1 = census.import_census(session, sheets)
    counts = (len(session.exec(select(Bed)).all()), len(session.exec(select(Assignment)).all()))
    assert r1.facilities_created == 3 and counts == (6 + 4 + 2, 5)
    r2 = census.import_census(session, sheets)
    assert (len(session.exec(select(Bed)).all()), len(session.exec(select(Assignment)).all())) == counts
    assert r2.facilities_created == 0
    assert sum(s.stays_created for s in r2.sheets) == 0 and sum(s.beds_created for s in r2.sheets) == 0
    bed = session.get(Bed, "rc-test-102-ada")
    assert bed.status == "needs_cleaning" and bed.wheelchair_accessible
    assert session.get(Bed, "rc-test-101-a").status == "occupied"
    assert session.get(Site, "rc-test").program_type == "recuperative_care"


def test_changed_occupant_closes_old_stay(session):
    census.import_census(session, parsed()[0])
    sheets, _ = parsed()
    row = next(r for r in by_sheet(sheets)["RC Test"].rows if r.label == "101-A")
    row.stay["client_id"], row.stay["first_name"] = "IF2099-77", "Newcomer"
    rep = census.import_census(session, sheets)
    rc = next(s for s in rep.sheets if s.sheet == "RC Test")
    assert rc.stays_closed == 1 and rc.stays_created == 1
    stays = session.exec(select(Assignment).where(Assignment.bed_id == "rc-test-101-a")).all()
    assert sorted(a.status for a in stays) == ["active", "released"]


def test_vacated_bed_closes_stay_and_frees_bed(session):
    census.import_census(session, parsed()[0])
    sheets, _ = parsed()
    row = next(r for r in by_sheet(sheets)["RC Test"].rows if r.label == "101-A")
    row.stay, row.status = None, "available"
    census.import_census(session, sheets)
    assert session.get(Bed, "rc-test-101-a").status == "available"


def test_referral_stay_preserved_and_conflict_reported(session):
    sheets, _ = parsed()
    census.import_census(session, sheets)
    bed = session.get(Bed, "rc-test-101-b")  # open in the sheet
    assert bed.status == "available"
    a = service.assign(session, bed.id, "ref-1", "Invented Person")
    # the sheet now shows someone in 101-B
    row = next(r for r in by_sheet(sheets)["RC Test"].rows if r.label == "101-B")
    row.status = "occupied"
    row.stay = {"client_id": "IF2099-9", "first_name": "Other", "last_name": "P.", "gender": None,
                "intake_date": None, "expected_exit_date": None, "exit_date": None, "payer": None,
                "case_manager": None, "notes": "", "extra": {}}
    rep = census.import_census(session, sheets)
    rc = next(s for s in rep.sheets if s.sheet == "RC Test")
    assert any("101-B" in c for c in rc.conflicts)
    session.expire_all()
    assert session.get(Assignment, a.id).status == "active"
    assert session.get(Bed, bed.id).status == "occupied"
    assert not session.exec(select(Assignment).where(Assignment.bed_id == bed.id, Assignment.source == "census")).all()


def test_existing_facility_county_not_overwritten(session):
    session.add(Site(id="rc-test", name="RC Test", county="Orange", program_type="shelter"))
    session.commit()
    census.import_census(session, parsed()[0])
    site = session.get(Site, "rc-test")
    assert site.county == "Orange" and site.program_type == "shelter"


def test_facility_program_hint_from_name(session):
    sheets, _ = parsed()
    census.import_census(session, sheets, program_hint=False)
    assert session.get(Site, "rc-test").program_type == "unknown"
    assert session.get(Site, "rc-test").county is None


def test_seed_poc_loads_placebo_only(session):
    n = seed_poc_beds(session)
    assert n >= 600
    assert len(session.exec(select(Site)).all()) >= 15
    assert seed_poc_beds(session) == 0  # only on an empty DB
    stays = session.exec(select(Assignment)).all()
    assert stays and all(a.client_id.startswith("POC-") for a in stays)
    assert all(a.first_name == "Test" and a.last_name.startswith("Client ") for a in stays)
    assert all(a.source == "seed" for a in stays)


def test_seed_json_contains_no_person_names():
    data = json.loads((Path(__file__).resolve().parent.parent / "app" / "beds" / "seed_poc.json").read_text("utf-8"))
    for f in data["facilities"]:
        assert f["notes"] == "POC placeholder county/program — confirm with IHH"
        for b in f["beds"]:
            st = b["stay"]
            if st:
                assert st["first_name"] == "Test" and st["last_name"].startswith("Client ")
                assert st["client_id"].startswith("POC-")
                assert st["extra"] == {"PETS": "No", "VI-SPDAT": "N/A", "Phone": "000-000-0000"}
                assert st["notes"] == "POC placebo occupant"


@pytest.mark.skipif(not REAL_WORKBOOK.exists(), reason="local synthetic census not present")
def test_real_workbook_parses():
    sheets, skipped = census.parse_workbook(REAL_WORKBOOK.read_bytes())
    sh = by_sheet(sheets)
    assert len(sheets) >= 18
    assert sum(len(s.rows) for s in sheets) >= 600
    for name in ("SFS", "Lincoln Park"):
        labels = [r.label for r in sh[name].rows]
        assert len(labels) >= 20 and len(labels) == len(set(labels))


def _workbook_with(builder) -> bytes:
    wb = openpyxl.Workbook()
    builder(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_block_without_if_column_is_kept_when_it_has_name_columns():
    """MWC-style: left block has no IF# column, right block does."""
    def build(wb):
        ws = wb.active
        ws.title = "MWC Test"
        ws.append(["Room #/Bed #", "/", "First Name", "Last Name", "Intake Date", "Notes", None, None, None,
                   "Room#/ Bed#", "IF#", "First Name", "Last ", "Intake"])
        ws.append(["1", None, "Open", None, None, None, None, None, None, "11A", None, "Open"])
        ws.append(["2", None, "Zed", "Q.", None, None, None, None, None, "11B", "IF2099-5", "Yan", "R."])
    sheets, _ = census.parse_workbook(_workbook_with(build))
    rows = {r.label: r for r in sheets[0].rows}
    assert set(rows) == {"1", "2", "11A", "11B"}
    assert rows["2"].stay["first_name"] == "Zed" and rows["11B"].stay["client_id"] == "IF2099-5"


def test_bed_attribute_column_does_not_split_a_block():
    def build(wb):
        ws = wb.active
        ws.title = "Attr Test"
        ws.append(["Room", "IF #", "First Name", "Last Name", "Intake Date", "Bed", "Pets", "Current Census"])
        ws.append(["1", "IF2099-6", "Xan", "S.", None, "DBL", "No"])
    sheets, _ = census.parse_workbook(_workbook_with(build))
    assert [r.label for r in sheets[0].rows] == ["1"]
    assert sheets[0].rows[0].stay["extra"]["Pets"] == "No"


def _seeded_bed(session):
    seed_poc_beds(session)
    occupied = session.exec(select(Assignment).where(Assignment.status == "active")).first()
    return session.get(Bed, occupied.bed_id)


def test_census_import_closes_seed_stay_when_bed_empty(session):
    bed = _seeded_bed(session)
    sheets, _ = parsed()
    sheets[0].sheet = session.get(Site, bed.site_id).name
    sheets[0].rows = [census.ParsedRow(label=bed.label, status="available")]
    census.import_census(session, sheets)
    session.expire_all()
    active = session.exec(select(Assignment).where(Assignment.bed_id == bed.id, Assignment.status == "active")).all()
    assert active == [] and session.get(Bed, bed.id).status == "available"


def test_census_import_replaces_seed_stay_with_new_occupant(session):
    bed = _seeded_bed(session)
    sheets, _ = parsed()
    sheets[0].sheet = session.get(Site, bed.site_id).name
    stay = {"client_id": "IF2099-8", "first_name": "Fresh", "last_name": "F.", "gender": None, "intake_date": None,
            "expected_exit_date": None, "exit_date": None, "payer": None, "case_manager": None, "notes": "",
            "extra": {}}
    sheets[0].rows = [census.ParsedRow(label=bed.label, status="occupied", stay=stay)]
    census.import_census(session, sheets)
    active = session.exec(select(Assignment).where(Assignment.bed_id == bed.id, Assignment.status == "active")).all()
    assert [a.first_name for a in active] == ["Fresh"]


@pytest.mark.skipif(not REAL_WORKBOOK.exists(), reason="local synthetic census not present")
def test_real_workbook_block_counts():
    sheets, _ = census.parse_workbook(REAL_WORKBOOK.read_bytes())
    sh = by_sheet(sheets)
    labels = [r.label for r in sh["MWC"].rows]
    assert "11A" in labels and "1" in labels and len(labels) >= 30
    assert len(sh["Royal Pagoda PRK"].rows) == 38
    assert len(sh["Laguna Hills VPS"].rows) == 78
    assert len(sh["Anaheim VPS"].rows) == 301

